#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Created on Mon Dec 13 15:25:58 2021.

@author: manab
"""
from openpyxl.drawing.colors import ColorChoice
from openpyxl.drawing.fill import PatternFillProperties


def _make_spPr_from_dict(target, props):
    """
    xlsxグラフのspPr(GraphicalProperties)を辞書型データから作成する関数.

    Parameters
    ----------
    target : openpyxl.chart.shapes
        spPrを指定する要素。
    props : DICT
        spPrの設定を指定する辞書型データ
    Returns
    -------
    None.
    """
    for key, value in props.items():
        if isinstance(value, dict):
            if key == "pattFill":
                fill = PatternFillProperties()
                for key_sub, value_sub in value.items():
                    if key_sub in ["foreground", "background"]:
                        setattr(fill, key_sub, ColorChoice(srgbClr=value_sub))
                    else:
                        setattr(fill, key_sub, value_sub)
                setattr(target.spPr, key, fill)
            elif key == "line":
                for key_sub, value_sub in value.items():
                    if key_sub == "width":
                        setattr(target.spPr.line, key_sub, value_sub * 12700)
                    else:
                        setattr(target.spPr.line, key_sub, value_sub)
            elif key == "marker":
                for key_sub, value_sub in value.items():
                    if key_sub == "spPr":
                        for key_sub2, value_sub2 in value_sub.items():
                            if key_sub2 == "line":
                                for key_sub3, value_sub3 in value_sub2.items():
                                    if key_sub3 == "width":
                                        setattr(
                                            target.marker.spPr.line, key_sub3, value_sub3 * 12700
                                        )
                                    else:
                                        setattr(target.marker.spPr.line, key_sub3, value_sub3)
                            else:
                                setattr(target.marker.spPr, key_sub2, value_sub2)
                    else:
                        setattr(target.marker, key_sub, value_sub)
        else:
            setattr(target.spPr, key, value)
