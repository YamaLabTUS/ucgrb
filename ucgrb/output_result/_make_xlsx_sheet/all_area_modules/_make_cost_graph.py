#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Cost グラフ生成モジュール.

全地域の目的関数（コスト）を示すグラフを作成する。
"""
from openpyxl.styles import Font

from .._append_col import _append_col
from .._make_object_function_chart import _make_object_function_chart


def make_cost_graph(ws, period_name, timeline, time_format, m, uc_data, uc_dicts):
    """
    全地域のコストグラフを作成する.

    Parameters
    ----------
    ws : CLASS
        結果を出力するシートのインスタンス
    period_name : STR
        表示対象の期間名称
    timeline : dataframe
        時系列
    time_format : STR
        時系列表示フォーマット
    m : CLASS
        Gurobiモデル
    uc_data : CLASS
        クラス「UCData」のインスタンス
    uc_dicts : CLASS
        クラス「UCDicts」のインスタンス

    """
    tsg_ratio = int(uc_data.config["time_series_granularity"]) / 60

    # 高速化のため、リストを事前に取得（一度だけ生成）
    area_list = list(uc_dicts.area)
    ess_list = list(uc_dicts.ess)
    tie_list = list(uc_dicts.tie)

    # 目的関数(コスト)
    ws.insert_cols(ws.max_column, 2)  # 総コスト表示用に2列追加
    _header_col = ["Cost"] + list(timeline.keys().strftime(time_format)) + ["Sum"]
    _append_col(ws, _header_col)
    _start_col = ws.max_column
    total = 0

    for g_type in uc_dicts.n_and_t_generation_type:
        _value_col = ["Coef (" + g_type + ")"]
        _sum = 0
        # 高速化のため、該当する発電機のリストを事前に取得
        filtered_generation_list = [
            (name, g_type_item, area_item)
            for name, g_type_item, area_item in uc_dicts.n_and_t_generation.select(
                "*", g_type, "*"
            )
        ]
        for time in timeline:
            if uc_data.config["optimization_timing"] == "intra-day" and hasattr(uc_dicts, "P_da"):
                # 当日計画の場合、P_daのみで構成
                _value = sum(
                    uc_dicts.P_da.get((time, name, g_type_item, area_item), 0)
                    * uc_dicts.generation_para["C_coef"][name, g_type_item, area_item]
                    * tsg_ratio
                    for name, g_type_item, area_item in filtered_generation_list
                    if (time, name, g_type_item, area_item) in uc_dicts.P_da
                )
                _value_col.append(_value)
                _sum += _value
            else:
                # 前日計画の場合、直接アクセスで高速化
                _value = sum(
                    uc_dicts.p[time, name, g_type_item, area_item].X
                    * uc_dicts.generation_para["C_coef"][name, g_type_item, area_item]
                    * tsg_ratio
                    for name, g_type_item, area_item in filtered_generation_list
                )
                _value_col.append(_value)
                _sum += _value
        _value_col.append(_sum)
        _append_col(ws, _value_col)
        total += _sum

        _value_col = ["Intc (" + g_type + ")"]
        _sum = 0
        # 高速化のため、該当する発電機のリストを事前に取得
        filtered_generation_list = [
            (name, g_type_item, area_item)
            for name, g_type_item, area_item in uc_dicts.n_and_t_generation.select(
                "*", g_type, "*"
            )
        ]
        for time in timeline:
            _value = sum(
                uc_dicts.u[time, name, g_type_item, area_item].X
                * uc_dicts.generation_para["C_intc"][name, g_type_item, area_item]
                * tsg_ratio
                for name, g_type_item, area_item in filtered_generation_list
            )
            _value_col.append(_value)
            _sum += _value
        _value_col.append(_sum)
        _append_col(ws, _value_col)
        total += _sum

        _value_col = ["Start Up (" + g_type + ")"]
        _sum = 0
        # 高速化のため、該当する発電機のリストを事前に取得
        filtered_generation_list = [
            (name, g_type_item, area_item)
            for name, g_type_item, area_item in uc_dicts.n_and_t_generation.select(
                "*", g_type, "*"
            )
        ]
        for time in timeline:
            _value = sum(
                uc_dicts.su[time, name, g_type_item, area_item].X
                * uc_dicts.generation_para["C_startup"][name, g_type_item, area_item]
                for name, g_type_item, area_item in filtered_generation_list
            )
            _value_col.append(_value)
            _sum += _value
        _value_col.append(_sum)
        _append_col(ws, _value_col)
        total += _sum

        # 大規模発電機の三次調整力コスト（ΔkW価値考慮版でのみ・前日/当日で分離）
        # delta-kW-no-market では ΔkW コスト内訳列を出さない
        if uc_data.config.get("formulation_type") != "delta-kW-bid":
            pass
        elif uc_data.config["optimization_timing"] == "day-ahead":
            # 前日計画: 三次調整力（ΔkW）のコストを発電機種類ごとに内訳表示
            if hasattr(uc_dicts, "p_delta_kW_tert_up") and hasattr(
                uc_dicts, "p_delta_kW_tert_down"
            ):
                # Tertiary Reserve Cost (Up) [ΔkW] を発電機種類ごとに内訳表示
                if g_type in uc_dicts.generation_type:
                    _value_col = ["Tertiary Reserve Cost (Up) [ΔkW] (" + g_type + ")"]
                    _sum = 0
                    # 高速化のため、該当する発電機のリストを事前に取得
                    filtered_generation_list = [
                        (name, g_type_item, area_item)
                        for name, g_type_item, area_item in uc_dicts.generation.select(
                            "*", g_type, "*"
                        )
                    ]
                    for time in timeline:
                        _value = sum(
                            uc_dicts.p_delta_kW_tert_up[time, name, g_type_item, area_item].X
                            * uc_dicts.generation_para["C_delta_kW_tert_UP"][
                                name, g_type_item, area_item
                            ]
                            * tsg_ratio
                            for name, g_type_item, area_item in filtered_generation_list
                        )
                        _value_col.append(_value)
                        _sum += _value
                    _value_col.append(_sum)
                    _append_col(ws, _value_col)
                    total += _sum

                    # Tertiary Reserve Cost (Down) [ΔkW] を発電機種類ごとに内訳表示
                    _value_col = ["Tertiary Reserve Cost (Down) [ΔkW] (" + g_type + ")"]
                    _sum = 0
                    # 高速化のため、該当する発電機のリストを事前に取得
                    filtered_generation_list = [
                        (name, g_type_item, area_item)
                        for name, g_type_item, area_item in uc_dicts.generation.select(
                            "*", g_type, "*"
                        )
                    ]
                    for time in timeline:
                        _value = sum(
                            uc_dicts.p_delta_kW_tert_down[time, name, g_type_item, area_item].X
                            * uc_dicts.generation_para["C_delta_kW_tert_DOWN"][
                                name, g_type_item, area_item
                            ]
                            * tsg_ratio
                            for name, g_type_item, area_item in filtered_generation_list
                        )
                        _value_col.append(_value)
                        _sum += _value
                    _value_col.append(_sum)
                    _append_col(ws, _value_col)
                    total += _sum
        elif uc_data.config["optimization_timing"] == "intra-day":
            # 当日計画: 三次調整電力量（kWh）のコストを発電機種類ごとに内訳表示
            if (
                hasattr(uc_dicts, "p_id_up")
                and hasattr(uc_dicts, "p_id_down")
                and "C_kWh_UP" in uc_dicts.generation_para
                and "C_kWh_DOWN" in uc_dicts.generation_para
            ):
                if g_type in uc_dicts.generation_type:
                    _value_col = ["Tertiary Energy Cost (Up) [kWh] (" + g_type + ")"]
                    _sum = 0
                    # 高速化のため、該当する発電機のリストを事前に取得
                    filtered_generation_list = [
                        (name, g_type_item, area_item)
                        for name, g_type_item, area_item in uc_dicts.generation.select(
                            "*", g_type, "*"
                        )
                    ]
                    for time in timeline:
                        _value = sum(
                            uc_dicts.p_id_up[time, name, g_type_item, area_item].X
                            * uc_dicts.generation_para["C_kWh_UP"].get(
                                (name, g_type_item, area_item), 0
                            )
                            * tsg_ratio
                            for name, g_type_item, area_item in filtered_generation_list
                        )
                        _value_col.append(_value)
                        _sum += _value
                    _value_col.append(_sum)
                    _append_col(ws, _value_col)
                    total += _sum

                    _value_col = ["Tertiary Energy Cost (Down) [kWh] (" + g_type + ")"]
                    _sum = 0
                    # 高速化のため、該当する発電機のリストを事前に取得
                    filtered_generation_list = [
                        (name, g_type_item, area_item)
                        for name, g_type_item, area_item in uc_dicts.generation.select(
                            "*", g_type, "*"
                        )
                    ]
                    for time in timeline:
                        _value = sum(
                            uc_dicts.p_id_down[time, name, g_type_item, area_item].X
                            * uc_dicts.generation_para["C_kWh_DOWN"].get(
                                (name, g_type_item, area_item), 0
                            )
                            * tsg_ratio
                            for name, g_type_item, area_item in filtered_generation_list
                        )
                        _value_col.append(_value)
                        _sum += _value
                    _value_col.append(_sum)
                    _append_col(ws, _value_col)
                    total += _sum

    # 水力発電機の三次調整力コストのみ表示（前日計画と当日計画で分離）
    for g_type in uc_dicts.hydro_generation_type:
        if uc_data.config.get("formulation_type") != "delta-kW-bid":
            pass  # delta-kW-no-market では ΔkW コスト内訳列を出さない
        elif uc_data.config["optimization_timing"] == "day-ahead":
            # 前日計画: 三次調整力（ΔkW）のコストを水力発電機種類ごとに内訳表示
            if hasattr(uc_dicts, "p_delta_kW_tert_up") and hasattr(
                uc_dicts, "p_delta_kW_tert_down"
            ):
                if g_type in uc_dicts.generation_type:
                    _value_col = ["Tertiary Reserve Cost (Up) [ΔkW] (" + g_type + ")"]
                    _sum = 0
                    # 高速化のため、該当する発電機のリストを事前に取得
                    filtered_generation_list = [
                        (name, g_type_item, area_item)
                        for name, g_type_item, area_item in uc_dicts.generation.select(
                            "*", g_type, "*"
                        )
                    ]
                    for time in timeline:
                        _value = sum(
                            uc_dicts.p_delta_kW_tert_up[time, name, g_type_item, area_item].X
                            * uc_dicts.generation_para["C_delta_kW_tert_UP"][
                                name, g_type_item, area_item
                            ]
                            * tsg_ratio
                            for name, g_type_item, area_item in filtered_generation_list
                        )
                        _value_col.append(_value)
                        _sum += _value
                    _value_col.append(_sum)
                    _append_col(ws, _value_col)
                    total += _sum

                    # Tertiary Reserve Cost (Down) [ΔkW] を水力発電機種類ごとに内訳表示
                    _value_col = ["Tertiary Reserve Cost (Down) [ΔkW] (" + g_type + ")"]
                    _sum = 0
                    # 高速化のため、該当する発電機のリストを事前に取得
                    filtered_generation_list = [
                        (name, g_type_item, area_item)
                        for name, g_type_item, area_item in uc_dicts.generation.select(
                            "*", g_type, "*"
                        )
                    ]
                    for time in timeline:
                        _value = sum(
                            uc_dicts.p_delta_kW_tert_down[time, name, g_type_item, area_item].X
                            * uc_dicts.generation_para["C_delta_kW_tert_DOWN"][
                                name, g_type_item, area_item
                            ]
                            * tsg_ratio
                            for name, g_type_item, area_item in filtered_generation_list
                        )
                        _value_col.append(_value)
                        _sum += _value
                    _value_col.append(_sum)
                    _append_col(ws, _value_col)
                    total += _sum
        elif uc_data.config["optimization_timing"] == "intra-day":
            # 当日計画: 三次調整電力量（kWh）のコストを水力発電機種類ごとに内訳表示
            if (
                hasattr(uc_dicts, "p_id_up")
                and hasattr(uc_dicts, "p_id_down")
                and "C_kWh_UP" in uc_dicts.generation_para
                and "C_kWh_DOWN" in uc_dicts.generation_para
            ):
                if g_type in uc_dicts.generation_type:
                    _value_col = ["Tertiary Energy Cost (Up) [kWh] (" + g_type + ")"]
                    _sum = 0
                    # 高速化のため、該当する発電機のリストを事前に取得
                    filtered_generation_list = [
                        (name, g_type_item, area_item)
                        for name, g_type_item, area_item in uc_dicts.generation.select(
                            "*", g_type, "*"
                        )
                    ]
                    for time in timeline:
                        _value = sum(
                            uc_dicts.p_id_up[time, name, g_type_item, area_item].X
                            * uc_dicts.generation_para["C_kWh_UP"].get(
                                (name, g_type_item, area_item), 0
                            )
                            * tsg_ratio
                            for name, g_type_item, area_item in filtered_generation_list
                        )
                        _value_col.append(_value)
                        _sum += _value
                    _value_col.append(_sum)
                    _append_col(ws, _value_col)
                    total += _sum

                    _value_col = ["Tertiary Energy Cost (Down) [kWh] (" + g_type + ")"]
                    _sum = 0
                    # 高速化のため、該当する発電機のリストを事前に取得
                    filtered_generation_list = [
                        (name, g_type_item, area_item)
                        for name, g_type_item, area_item in uc_dicts.generation.select(
                            "*", g_type, "*"
                        )
                    ]
                    for time in timeline:
                        _value = sum(
                            uc_dicts.p_id_down[time, name, g_type_item, area_item].X
                            * uc_dicts.generation_para["C_kWh_DOWN"].get(
                                (name, g_type_item, area_item), 0
                            )
                            * tsg_ratio
                            for name, g_type_item, area_item in filtered_generation_list
                        )
                        _value_col.append(_value)
                        _sum += _value
                    _value_col.append(_sum)
                    _append_col(ws, _value_col)
                    total += _sum

    _value_col = ["Short"]
    _sum = 0
    for time in timeline:
        _value = sum(
            uc_dicts.p_short[time, area_item].X
            * uc_dicts.area_para["C_short"][area_item]
            * tsg_ratio
            for area_item in area_list
        )
        _value_col.append(_value)
        _sum += _value
    _value_col.append(_sum)
    _append_col(ws, _value_col)
    total += _sum

    _value_col = ["Surplus"]
    _sum = 0
    for time in timeline:
        _value = sum(
            uc_dicts.p_surplus[time, area_item].X
            * uc_dicts.area_para["C_surplus"][area_item]
            * tsg_ratio
            for area_item in area_list
        )
        _value_col.append(_value)
        _sum += _value
    _value_col.append(_sum)
    _append_col(ws, _value_col)
    total += _sum

    _value_col = ["Suppression of PV"]
    _sum = 0
    for time in timeline:
        _value = sum(
            uc_dicts.p_pv_suppr[time, area_item].X
            * uc_dicts.area_para["C_PV_suppr"][area_item]
            * tsg_ratio
            for area_item in area_list
        )
        _value_col.append(_value)
        _sum += _value
    _value_col.append(_sum)
    _append_col(ws, _value_col)
    total += _sum

    _value_col = ["Suppression of WF"]
    _sum = 0
    for time in timeline:
        _value = sum(
            uc_dicts.p_wf_suppr[time, area_item].X
            * uc_dicts.area_para["C_WF_suppr"][area_item]
            * tsg_ratio
            for area_item in area_list
        )
        _value_col.append(_value)
        _sum += _value
    _value_col.append(_sum)
    _append_col(ws, _value_col)
    total += _sum

    _value_col = ["Short of Tertiary (UP)"]
    _sum = 0
    for time in timeline:
        _value = sum(
            uc_dicts.p_tert_up_short[time, area_item].X
            * uc_dicts.area_para["C_Tert_short"][area_item]
            * tsg_ratio
            for area_item in area_list
        )
        _value_col.append(_value)
        _sum += _value
    _value_col.append(_sum)
    _append_col(ws, _value_col)
    total += _sum

    _value_col = ["Short of Tertiary (DOWN)"]
    _sum = 0
    for time in timeline:
        _value = sum(
            uc_dicts.p_tert_down_short[time, area_item].X
            * uc_dicts.area_para["C_Tert_short"][area_item]
            * tsg_ratio
            for area_item in area_list
        )
        _value_col.append(_value)
        _sum += _value
    _value_col.append(_sum)
    _append_col(ws, _value_col)
    total += _sum

    _value_col = ["ESS Short Penalty"]
    _sum = 0
    for time in timeline:
        _value = sum(
            uc_dicts.e_ess_short[time, ess, area].X
            * uc_dicts.ess_para["C_ess_short"][ess, area]
            * tsg_ratio
            for ess, area in ess_list
        )
        _value_col.append(_value)
        _sum += _value
    _value_col.append(_sum)
    _append_col(ws, _value_col)
    total += _sum

    _value_col = ["ESS Surplus Penalty"]
    _sum = 0
    for time in timeline:
        _value = sum(
            uc_dicts.e_ess_surplus[time, ess, area].X
            * uc_dicts.ess_para["C_ess_surplus"][ess, area]
            * tsg_ratio
            for ess, area in ess_list
        )
        _value_col.append(_value)
        _sum += _value
    _value_col.append(_sum)
    _append_col(ws, _value_col)
    total += _sum

    # ESSの三次調整力コスト（前日計画と当日計画で分離）
    if uc_data.config.get("formulation_type") != "delta-kW-bid":
        pass  # delta-kW-no-market では ΔkW コスト内訳列を出さない
    elif uc_data.config["optimization_timing"] == "day-ahead":
        # 前日計画: 三次調整力（ΔkW）のコストを全ESS合計で表示
        if hasattr(uc_dicts, "p_ess_delta_kW_tert_up") and hasattr(
            uc_dicts, "p_ess_delta_kW_tert_down"
        ):
            # Tertiary Reserve Cost (Up) [ΔkW] を全ESS合計で表示
            _value_col = ["ESS Tertiary Reserve Cost (Up) [ΔkW]"]
            _sum = 0
            for time in timeline:
                _value = sum(
                    uc_dicts.p_ess_delta_kW_tert_up[time, name, area].X
                    * uc_dicts.ess_para["C_ess_delta_kW_tert_UP"][name, area]
                    * tsg_ratio
                    for name, area in ess_list
                )
                _value_col.append(_value)
                _sum += _value
            _value_col.append(_sum)
            _append_col(ws, _value_col)
            total += _sum

            # Tertiary Reserve Cost (Down) [ΔkW] を全ESS合計で表示
            _value_col = ["ESS Tertiary Reserve Cost (Down) [ΔkW]"]
            _sum = 0
            for time in timeline:
                _value = sum(
                    uc_dicts.p_ess_delta_kW_tert_down[time, name, area].X
                    * uc_dicts.ess_para["C_ess_delta_kW_tert_DOWN"][name, area]
                    * tsg_ratio
                    for name, area in ess_list
                )
                _value_col.append(_value)
                _sum += _value
            _value_col.append(_sum)
            _append_col(ws, _value_col)
            total += _sum
    elif uc_data.config["optimization_timing"] == "intra-day":
        # 当日計画: 三次調整電力量（kWh）のコスト
        if (
            hasattr(uc_dicts, "p_ess_id_up")
            and hasattr(uc_dicts, "p_ess_id_down")
            and "C_ess_kWh_UP" in uc_dicts.ess_para
            and "C_ess_kWh_DOWN" in uc_dicts.ess_para
        ):
            _value_col = ["ESS Tertiary Energy Cost (Up) [kWh]"]
            _sum = 0
            for time in timeline:
                _value = sum(
                    uc_dicts.p_ess_id_up[time, name, area].X
                    * uc_dicts.ess_para["C_ess_kWh_UP"].get((name, area), 0)
                    * tsg_ratio
                    for name, area in ess_list
                )
                _value_col.append(_value)
                _sum += _value
            _value_col.append(_sum)
            _append_col(ws, _value_col)
            total += _sum

            _value_col = ["ESS Tertiary Energy Cost (Down) [kWh]"]
            _sum = 0
            for time in timeline:
                _value = sum(
                    uc_dicts.p_ess_id_down[time, name, area].X
                    * uc_dicts.ess_para["C_ess_kWh_DOWN"].get((name, area), 0)
                    * tsg_ratio
                    for name, area in ess_list
                )
                _value_col.append(_value)
                _sum += _value
            _value_col.append(_sum)
            _append_col(ws, _value_col)
            total += _sum

    total_wo_tie = total
    _value_col = ["Tie Line Used Penalty"]
    _sum = 0
    for time in timeline:
        _value = sum(
            uc_dicts.tie_para["C_tie_penalty"][name, f, t]
            * (uc_dicts.p_tie_f[time, name, f, t].X + uc_dicts.p_tie_c[time, name, f, t].X)
            * tsg_ratio
            for name, f, t in tie_list
        )
        _value_col.append(_value)
        _sum += _value
    _value_col.append(_sum)
    _append_col(ws, _value_col)
    total += _sum

    _value_col = ["Tie Line Used Penalty (GF&LFC, Up)"]
    _sum = 0
    for time in timeline:
        _value = sum(
            uc_dicts.tie_para["C_tie_penalty_GF_LFC_UP"][name, f, t]
            * (
                uc_dicts.p_tie_gf_lfc_up_f[time, name, f, t].X
                + uc_dicts.p_tie_gf_lfc_up_c[time, name, f, t].X
            )
            * tsg_ratio
            for name, f, t in tie_list
        )
        _value_col.append(_value)
        _sum += _value
    _value_col.append(_sum)
    _append_col(ws, _value_col)
    total += _sum

    _value_col = ["Tie Line Used Penalty (GF&LFC, Down)"]
    _sum = 0
    for time in timeline:
        _value = sum(
            uc_dicts.tie_para["C_tie_penalty_GF_LFC_DOWN"][name, f, t]
            * (
                uc_dicts.p_tie_gf_lfc_down_f[time, name, f, t].X
                + uc_dicts.p_tie_gf_lfc_down_c[time, name, f, t].X
            )
            * tsg_ratio
            for name, f, t in tie_list
        )
        _value_col.append(_value)
        _sum += _value
    _value_col.append(_sum)
    _append_col(ws, _value_col)
    total += _sum

    _value_col = ["Tie Line Used Penalty (Tert, Up)"]
    _sum = 0
    for time in timeline:
        _value = sum(
            uc_dicts.tie_para["C_tie_penalty_Tert_UP"][name, f, t]
            * (
                uc_dicts.p_tie_tert_up_f[time, name, f, t].X
                + uc_dicts.p_tie_tert_up_c[time, name, f, t].X
            )
            * tsg_ratio
            for name, f, t in tie_list
        )
        _value_col.append(_value)
        _sum += _value
    _value_col.append(_sum)
    _append_col(ws, _value_col)
    total += _sum

    _value_col = ["Tie Line Used Penalty (Tert, Down)"]
    _sum = 0
    for time in timeline:
        _value = sum(
            uc_dicts.tie_para["C_tie_penalty_Tert_DOWN"][name, f, t]
            * (
                uc_dicts.p_tie_tert_down_f[time, name, f, t].X
                + uc_dicts.p_tie_tert_down_c[time, name, f, t].X
            )
            * tsg_ratio
            for name, f, t in tie_list
        )
        _value_col.append(_value)
        _sum += _value
    _value_col.append(_sum)
    _append_col(ws, _value_col)
    total += _sum

    ws.cell(column=_start_col - 2, row=1, value="Total Cost [kJPY]")
    ws.cell(column=_start_col - 2, row=2, value=total_wo_tie)
    ws.cell(
        column=_start_col - 2,
        row=4,
        value="Total Cost with Used Penalty of Tie Line [kJPY]",
    )
    ws.cell(column=_start_col - 2, row=5, value=total)

    cordinate = ws.cell(column=_start_col - 2, row=2).coordinate
    ws[cordinate].font = Font(bold=True)

    _make_object_function_chart(
        ws,
        "Cost in All Area on " + period_name,
        place_row=2,
        place_col=ws.max_column + 2,
        start_row=1,
        start_col=_start_col,
        len_timeline=len(timeline),
        len_elements=len(uc_dicts.n_and_t_generation_type) * 3
        + (
            len(
                [
                    g_type
                    for g_type in uc_dicts.n_and_t_generation_type
                    if g_type in uc_dicts.generation_type
                ]
            )
            * 2
            if uc_data.config.get("formulation_type") == "delta-kW-bid"
            and uc_data.config["optimization_timing"] == "day-ahead"
            and hasattr(uc_dicts, "p_delta_kW_tert_up")
            else 0
        )
        + (
            len(
                [
                    g_type
                    for g_type in uc_dicts.n_and_t_generation_type
                    if g_type in uc_dicts.generation_type
                ]
            )
            * 2
            if uc_data.config.get("formulation_type") == "delta-kW-bid"
            and uc_data.config["optimization_timing"] == "intra-day"
            and hasattr(uc_dicts, "p_id_up")
            else 0
        )
        + (
            len(
                [
                    g_type
                    for g_type in uc_dicts.hydro_generation_type
                    if g_type in uc_dicts.generation_type
                ]
            )
            * 2
            if uc_data.config.get("formulation_type") == "delta-kW-bid"
            and uc_data.config["optimization_timing"] == "day-ahead"
            and hasattr(uc_dicts, "p_delta_kW_tert_up")
            else 0
        )
        + (
            len(
                [
                    g_type
                    for g_type in uc_dicts.hydro_generation_type
                    if g_type in uc_dicts.generation_type
                ]
            )
            * 2
            if uc_data.config.get("formulation_type") == "delta-kW-bid"
            and uc_data.config["optimization_timing"] == "intra-day"
            and hasattr(uc_dicts, "p_id_up")
            else 0
        )
        + 13
        + (
            2
            if uc_data.config.get("formulation_type") == "delta-kW-bid"
            and uc_data.config["optimization_timing"] == "day-ahead"
            and hasattr(uc_dicts, "p_ess_delta_kW_tert_up")
            else 0
        )
        + (
            2
            if uc_data.config.get("formulation_type") == "delta-kW-bid"
            and uc_data.config["optimization_timing"] == "intra-day"
            and hasattr(uc_dicts, "p_ess_id_up")
            else 0
        ),
        size=1.5,
        graphical_prop=uc_data.config["graphical_prop_for_xlsx_graph"],
    )

    ws.cell(column=ws.max_column + 1, row=1, value="")
    ws.insert_cols(ws.max_column, 18)
