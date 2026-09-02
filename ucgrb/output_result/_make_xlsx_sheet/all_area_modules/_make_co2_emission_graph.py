#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
CO2 Emission グラフ生成モジュール.

全地域のCO2排出量を示すグラフを作成する。
"""
from openpyxl.styles import Font

from .._append_col import _append_col
from .._make_object_function_chart import _make_object_function_chart


def make_co2_emission_graph(ws, period_name, timeline, time_format, m, uc_data, uc_dicts):
    """
    全地域のCO2排出量グラフを作成する.

    Parameters
    ----------
    ws : CLASS
        結果を出力するシートのインスタンス
    period_name : STR
        表示対象の期間名称
    timeline : dataframe
        時系列
    time_format : STR
        時系列表示フォーマット
    m : CLASS
        Gurobiモデル
    uc_data : CLASS
        クラス「UCData」のインスタンス
    uc_dicts : CLASS
        クラス「UCDicts」のインスタンス

    """
    tsg_ratio = int(uc_data.config["time_series_granularity"]) / 60

    # CO2排出量
    ws.insert_cols(ws.max_column, 2)  # 総排出量表示用に2列追加
    _header_col = ["CO2 Emission"] + list(timeline.keys().strftime(time_format)) + ["Sum"]
    _append_col(ws, _header_col)
    _start_col = ws.max_column
    total = 0

    num_g_type = 0
    for g_type in uc_dicts.generation_type:
        if g_type in ["HYDRO", "NUCL"]:
            continue
        num_g_type = num_g_type + 1
        _value_col = ["Coef (" + g_type + ")"]
        _sum = 0
        # 高速化のため、該当する発電機のリストを事前に取得
        filtered_generation_list = [
            (name, g_type_item, area_item)
            for name, g_type_item, area_item in uc_dicts.generation.select("*", g_type, "*")
        ]
        for time in timeline:
            # 直接アクセス + Pythonのsum()を使用（高速化）
            _value = sum(
                uc_dicts.p[time, name, g_type_item, area_item].X
                * uc_dicts.generation_para["C_coef_CO2"][name, g_type_item, area_item]
                * tsg_ratio
                for name, g_type_item, area_item in filtered_generation_list
            )
            _value_col.append(_value)
            _sum += _value
        _value_col.append(_sum)
        _append_col(ws, _value_col)
        total += _sum

        _value_col = ["Intc (" + g_type + ")"]
        _sum = 0
        # 高速化のため、該当する発電機のリストを事前に取得
        filtered_generation_list = [
            (name, g_type_item, area_item)
            for name, g_type_item, area_item in uc_dicts.n_and_t_generation.select(
                "*", g_type, "*"
            )
        ]
        for time in timeline:
            # 直接アクセス + Pythonのsum()を使用（高速化）
            _value = sum(
                uc_dicts.u[time, name, g_type_item, area_item].X
                * uc_dicts.generation_para["C_intc_CO2"][name, g_type_item, area_item]
                * tsg_ratio
                for name, g_type_item, area_item in filtered_generation_list
            )
            _value_col.append(_value)
            _sum += _value
        _value_col.append(_sum)
        _append_col(ws, _value_col)
        total += _sum

        _value_col = ["Start Up (" + g_type + ")"]
        _sum = 0
        # 高速化のため、該当する発電機のリストを事前に取得
        filtered_generation_list = [
            (name, g_type_item, area_item)
            for name, g_type_item, area_item in uc_dicts.n_and_t_generation.select(
                "*", g_type, "*"
            )
        ]
        for time in timeline:
            # 直接アクセス + Pythonのsum()を使用（高速化）
            _value = sum(
                uc_dicts.su[time, name, g_type_item, area_item].X
                * uc_dicts.generation_para["C_startup_CO2"][name, g_type_item, area_item]
                for name, g_type_item, area_item in filtered_generation_list
            )
            _value_col.append(_value)
            _sum += _value
        _value_col.append(_sum)
        _append_col(ws, _value_col)
        total += _sum

    ws.cell(column=_start_col - 2, row=1, value="Total CO2 Emission [tCO2]")
    ws.cell(column=_start_col - 2, row=2, value=total)

    cordinate = ws.cell(column=_start_col - 2, row=2).coordinate
    ws[cordinate].font = Font(bold=True)

    _make_object_function_chart(
        ws,
        "CO2 Emission in All Area on " + period_name,
        place_row=2,
        place_col=ws.max_column + 2,
        start_row=1,
        start_col=_start_col,
        len_timeline=len(timeline),
        len_elements=num_g_type * 3,
        y_axis_title="[tCO2]",
        graphical_prop=uc_data.config["graphical_prop_for_xlsx_graph"],
    )

    ws.cell(column=ws.max_column + 1, row=1, value="")
    ws.insert_cols(ws.max_column, 12)
